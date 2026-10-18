"""Build the approved Cinema SRS as a styled, navigable Word document.

Dependencies: python-docx, pymupdf; LibreOffice for PDF rendering/QA.
Run once, render to PDF, then run with --pages <rendered.pdf> to fill TOC numbers.
The source of truth remains ../../SRS.md. No production credentials are read.
"""
from pathlib import Path
import argparse
import re
import unicodedata
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT.parents[1] / 'SRS.md'
OUTPUT = ROOT / 'SRS_Cinema_v1.0.docx'
NAVY, TEAL, GRAY = '17324D', '007F86', '526477'
WIDTH = 16.6


def element(tag, **attrs):
    e = OxmlElement('w:' + tag)
    for key, val in attrs.items():
        e.set(qn('w:' + key), str(val))
    return e


def shade(parent, fill):
    parent.append(element('shd', fill=fill, val='clear'))


def rich(p, text, color=None, size=None):
    for fragment in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not fragment:
            continue
        bold = fragment.startswith('**') and fragment.endswith('**')
        code = fragment.startswith('`') and fragment.endswith('`')
        text = fragment[2:-2] if bold else fragment[1:-1] if code else fragment
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if code:
            r.font.name = 'Consolas'
            r.font.size = Pt(min(size or 10, 9))
    return p


def bookmark(p, ident):
    p._p.append(element('bookmarkStart', id=ident, name=f'srs_{ident}'))
    p._p.append(element('bookmarkEnd', id=ident))


def field(p, instr):
    r = p.add_run()
    r._r.append(element('fldChar', fldCharType='begin'))
    e = element('instrText')
    e.text = ' ' + instr + ' '
    r._r.append(e)
    r._r.append(element('fldChar', fldCharType='end'))


def toc_entry(doc, title, ident, level, page):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.0
    pf.left_indent = Cm(0.35 if level == 2 else 0)
    pf.tab_stops.add_tab_stop(Cm(WIDTH), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    link = element('hyperlink', anchor=f'srs_{ident}', history='1')
    run = element('r')
    prop = element('rPr')
    prop.append(element('color', val=NAVY))
    prop.append(element('sz', val='20'))
    if level == 1:
        prop.append(element('b'))
    run.append(prop)
    txt = element('t')
    txt.text = title
    run.append(txt)
    link.append(run)
    p._p.append(link)
    p.add_run('\t' + str(page or '—')).font.size = Pt(10)


def widths_for(headers):
    if headers[0] == 'Phiên bản':
        return [1.4, 2.1, 9.0, 4.1]
    if headers[0] == 'Yêu cầu' and len(headers) == 4:
        return [3.5, 2.1, 4.0, 7.0]
    if len(headers) == 4:
        return [1.6, 5.1, 4.5, 5.4]
    if len(headers) == 3:
        if headers[0] == 'Case':
            return [3.5, 6.7, 6.4]
        if headers[1] == 'Nội dung':
            return [2.0, 5.0, 9.6]
        return [3.3, 8.1, 5.2]
    if headers[0] == 'Mã':
        return [3.1, 13.5]
    return [4.1, 12.5]


def add_table(doc, rows):
    widths = widths_for(rows[0])
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    props = table._tbl.tblPr
    borders = element('tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        borders.append(element(side, val='single', sz='4', color='D9E2EA'))
    props.append(borders)
    margins = element('tblCellMar')
    for side in ['top', 'bottom', 'left', 'right']:
        margins.append(element(side, w='60' if side in ('top', 'bottom') else '90', type='dxa'))
    props.append(margins)
    for idx, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(element('cantSplit'))
        if idx == 0:
            row._tr.get_or_add_trPr().append(element('tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            shade(cell._tc.get_or_add_tcPr(), NAVY if idx == 0 else ('F0F5F8' if idx % 2 else 'FFFFFF'))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.06
            rich(p, value, 'FFFFFF' if idx == 0 else None, 9.3)
            if idx == 0:
                for r in p.runs:
                    r.bold = True
                p.paragraph_format.keep_with_next = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.2
    spacer.add_run().font.size = Pt(2)


def normalized(text):
    return re.sub(r'\s+', '', unicodedata.normalize('NFC', text)).casefold()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--pages', type=Path)
    args = parser.parse_args()
    lines = SOURCE.read_text().splitlines()
    start = next(i for i, x in enumerate(lines) if x == '## Kiểm soát tài liệu')
    body = lines[start:]
    headings = []
    for line in body:
        m = re.match(r'^(#{1,2}) (.+)$', line)
        if m:
            level = 1 if len(m[1]) == 1 or m[2].startswith('Phụ lục') else 2
            headings.append((m[2], level))
    ids = {title: idx + 1 for idx, (title, _) in enumerate(headings)}
    pages = {}
    if args.pages:
        import fitz
        with fitz.open(args.pages) as pdf:
            # Cover and the two deliberately fixed TOC pages are front matter.
            for n in range(3, len(pdf)):
                text = normalized(pdf[n].get_text())
                for title, _ in headings:
                    if title not in pages and normalized(title) in text:
                        pages[title] = n + 1
        if len(pages) != len(headings):
            raise RuntimeError('Missing headings in rendered PDF: ' + str(set(ids) - set(pages)))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.9)
    sec.left_margin, sec.right_margin = Cm(2.4), Cm(2.0)
    sec.header_distance = sec.footer_distance = Cm(0.8)
    sec.different_first_page_header_footer = True
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 28), ('Subtitle', 14), ('Heading 1', 18), ('Heading 2', 13), ('Heading 3', 11.5)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Arial', Pt(size)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.font.bold = name != 'Subtitle'
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
    for name in ['List Bullet', 'List Number']:
        doc.styles[name].font.name = 'Arial'
        doc.styles[name].font.size = Pt(10.5)
    header = sec.header.paragraphs[0]
    rich(header, 'CINEMA  /  ĐẶC TẢ YÊU CẦU PHẦN MỀM', GRAY, 8)
    header.paragraph_format.space_after = Pt(5)
    borders = element('pBdr')
    borders.append(element('bottom', val='single', sz='6', color=TEAL, space='5'))
    header._p.get_or_add_pPr().append(borders)
    footer = sec.footer.paragraphs[0]
    footer.style = doc.styles['Normal']
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    rich(footer, 'CINEMA-SRS-1.0  •  Baseline đã phê duyệt\tTrang ', GRAY, 8)
    field(footer, 'PAGE')
    footer.add_run(' / ').font.size = Pt(8)
    field(footer, 'NUMPAGES')

    # Cover: no invented institution, author name or signatures.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(55)
    rich(p, 'CINEMA', TEAL, 16).runs[0].bold = True
    doc.add_paragraph('ĐẶC TẢ YÊU CẦU\nPHẦN MỀM', 'Title')
    doc.add_paragraph('Software Requirements Specification', 'Subtitle')
    p = doc.add_paragraph('Hệ thống quản lý rạp chiếu phim\nvà đặt vé trực tuyến')
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(32)
    for r in p.runs:
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor.from_string(NAVY)
    add_table(doc, [
        ['Thông tin tài liệu', 'Nội dung'],
        ['Mã tài liệu', 'CINEMA-SRS-1.0'],
        ['Phiên bản', '1.0 — Baseline đã phê duyệt'],
        ['Ngày phê duyệt', '18/10/2026'],
        ['Người phê duyệt', 'Chủ dự án — xác nhận trong cuộc trao đổi'],
        ['Mục đích', 'Đồ án phần mềm; demo và kiểm thử Stripe Sandbox'],
        ['Ngôn ngữ', 'Tiếng Việt'],
    ])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    rich(p, 'Phê duyệt yêu cầu không thay thế nghiệm thu hệ thống.\nCác giới hạn kiểm thử và phạm vi Sandbox được nêu rõ trong tài liệu.', GRAY, 10)
    doc.add_page_break()

    # Fixed two-page TOC keeps page references stable when they are filled in.
    midpoint = (len(headings) + 1) // 2
    for part, items in enumerate([headings[:midpoint], headings[midpoint:]]):
        doc.add_paragraph('MỤC LỤC' + (' (tiếp theo)' if part else ''), 'Title')
        if not part:
            rich(doc.add_paragraph(), 'Chọn tên mục để chuyển đến nội dung tương ứng. Số trang tham chiếu bản phát hành v1.0.', GRAY, 9)
        for title, level in items:
            toc_entry(doc, title, ids[title], level, pages.get(title))
        doc.add_page_break()

    i = 0
    while i < len(body):
        line = body[i].strip()
        if not line or line == '---':
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(body) and body[i].strip().startswith('|'):
                cells = [v.strip() for v in body[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v) for v in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith('```'):
            code = []
            i += 1
            while i < len(body) and not body[i].startswith('```'):
                code.append(body[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.12
            shade(p._p.get_or_add_pPr(), 'F0F5F8')
            for n, text in enumerate(code):
                r = p.add_run(('\n' if n else '') + text)
                r.font.name, r.font.size = 'Consolas', Pt(9)
            i += 1
            continue
        match = re.match(r'^(#{1,3}) (.+)$', line)
        if match:
            level, title = len(match[1]), match[2]
            if title.startswith('Phụ lục'):
                level = 1
            p = doc.add_paragraph(title, 'Heading ' + str(level))
            if (level == 1 and not title.startswith('Phụ lục B')) or title in {
                '1.4. Thuật ngữ', '2.4. Môi trường và phụ thuộc',
                '3.5. Quy tắc nghiệp vụ', '5.3. Ràng buộc toàn vẹn',
                '6.6. Ma trận nghiệm thu Stripe Sandbox', '6.7. Truy vết yêu cầu',
            }:
                p.paragraph_format.page_break_before = True
            if title in ids:
                bookmark(p, ids[title])
        elif line.startswith('> '):
            p = rich(doc.add_paragraph(), line[2:], GRAY, 10)
            shade(p._p.get_or_add_pPr(), 'EAF4F4')
        elif line.startswith('- '):
            rich(doc.add_paragraph(style='List Bullet'), line[2:])
        elif re.match(r'^\d+\. ', line):
            # Explicit numbering preserves independent numbered lists from the source.
            p = rich(doc.add_paragraph(), line)
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.55)
        else:
            rich(doc.add_paragraph(), line)
        i += 1
    props = doc.core_properties
    props.title = 'Cinema — Đặc tả yêu cầu phần mềm'
    props.subject = 'SRS baseline v1.0 đã được chủ dự án phê duyệt'
    props.author = 'Dự án Cinema'
    props.keywords = 'Cinema, SRS, Stripe Sandbox, baseline 1.0'
    props.comments = 'Nguồn: SRS.md. Phê duyệt yêu cầu không đồng nghĩa nghiệm thu toàn bộ hệ thống.'
    props.version = '1.0'
    doc.save(OUTPUT)
    print(f'Created {OUTPUT}; {len(headings)} TOC entries, {len(doc.tables)} tables.')


if __name__ == '__main__':
    main()
